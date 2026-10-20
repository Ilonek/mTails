"""Write the mTails JSON model to an Excel workbook (openpyxl)."""

from __future__ import annotations

import base64
import io
from typing import Any, Dict, Optional

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.protection import SheetProtection
from openpyxl.drawing.image import Image as XLImage
from openpyxl.workbook.defined_name import DefinedName

from .formats import default_page

FONT_FALLBACK = "Calibri"

BORDER_SIDES = {
    "thin": "thin",
    "medium": "medium",
    "thick": "thick",
    "dashed": "dashed",
    "dotted": "dotted",
    "double": "double",
    "hair": "hair",
    "mediumDashed": "mediumDashed",
    "dashDot": "dashDot",
    "mediumDashDot": "mediumDashDot",
    "dashDotDot": "dashDotDot",
}


def _font(style: dict) -> Font:
    color = style.get("color")
    underline = style.get("u")
    if underline is True:
        underline = "single"
    return Font(
        name=style.get("font") or FONT_FALLBACK,
        size=float(style.get("size") or 11),
        bold=bool(style.get("b")),
        italic=bool(style.get("i")),
        underline=underline or None,
        strike=bool(style.get("strike")),
        color=(color[1:] if isinstance(color, str) and color.startswith("#") else color) or None,
    )


def _fill(style: dict) -> Optional[PatternFill]:
    color = style.get("fill")
    if not color:
        return None
    value = color[1:] if color.startswith("#") else color
    return PatternFill(patternType="solid", fgColor=value)


def _border(style: dict) -> Optional[Border]:
    spec = style.get("border")
    if not spec:
        return None
    kwargs = {}
    for key, name in (("t", "top"), ("r", "right"), ("b", "bottom"), ("l", "left")):
        entry = spec.get(key)
        if not entry:
            continue
        color = entry.get("color")
        kwargs[name] = Side(
            style=BORDER_SIDES.get(entry.get("style", "thin"), "thin"),
            color=(color[1:] if isinstance(color, str) and color.startswith("#") else color) or None,
        )
    return Border(**kwargs) if kwargs else None


def _alignment(style: dict) -> Optional[Alignment]:
    horizontal = style.get("align")
    vertical = style.get("valign")
    wrap = bool(style.get("wrap"))
    indent = int(style.get("indent") or 0)
    if not horizontal and not vertical and not wrap and not indent:
        return None
    return Alignment(
        horizontal=horizontal,
        vertical={"middle": "center"}.get(vertical, vertical),
        wrap_text=wrap,
        indent=indent,
    )


def _apply_style(cell, style: dict, style_cache: Dict[str, Any]) -> None:
    if not style:
        return
    key = repr(sorted(style.items(), key=lambda item: item[0]))
    cached = style_cache.get(key)
    if cached is None:
        cached = (_font(style), _fill(style), _border(style), _alignment(style), style.get("numFmt"))
        style_cache[key] = cached
    font, fill, border, alignment, number_format = cached
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if border is not None:
        cell.border = border
    if alignment is not None:
        cell.alignment = alignment
    if number_format:
        cell.number_format = number_format


def _value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        # Excel limits a cell to 32767 characters.
        return value[:32767]
    return str(value)


def _drawing_anchor(sheet: dict, x: int, y: int, width: int, height: int) -> OneCellAnchor:
    """Place a drawing at a pixel offset using a one-cell anchor."""
    col = 1
    left = 0
    while col < 1000:
        size = int((sheet.get("cols", {}).get(get_column_letter(col)) or {}).get("width") or 96)
        if left + size > x:
            break
        left += size
        col += 1
    row = 1
    top = 0
    while row < 100000:
        size = int((sheet.get("rows", {}).get(str(row)) or {}).get("height") or 22)
        if top + size > y:
            break
        top += size
        row += 1
    marker = AnchorMarker(
        col=max(0, col - 1),
        colOff=pixels_to_EMU(max(0, x - left)),
        row=max(0, row - 1),
        rowOff=pixels_to_EMU(max(0, y - top)),
    )
    size = XDRPositiveSize2D(pixels_to_EMU(max(1, width)), pixels_to_EMU(max(1, height)))
    return OneCellAnchor(_from=marker, ext=size)


def _chart(sheet: dict, spec: dict, ws) -> Optional[Any]:
    range_text = str(spec.get("range") or "")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_text.replace("$", ""))
    except Exception:
        return None
    if min_col is None:
        return None
    kind = spec.get("type") or "column"
    if kind == "bar":
        chart = BarChart()
        chart.type = "bar"
    elif kind in ("column", "sparkline"):
        chart = BarChart()
        chart.type = "col"
    elif kind == "line":
        chart = LineChart()
    elif kind == "pie":
        chart = PieChart()
    elif kind == "area":
        chart = AreaChart()
    elif kind == "scatter":
        chart = ScatterChart()
    else:
        chart = BarChart()
    header_row = bool(spec.get("headerRow", True))
    header_col = bool(spec.get("headerCol", False))
    data_min_row = min_row + 1 if header_row and max_row > min_row else min_row
    data = Reference(ws, min_col=min_col, min_row=data_min_row, max_col=max_col, max_row=max_row)
    categories = Reference(
        ws,
        min_col=min_col,
        min_row=data_min_row,
        max_col=min_col,
        max_row=max_row,
    )
    chart.add_data(data, titles_from_data=header_col)
    chart.set_categories(categories)
    if spec.get("title"):
        chart.title = str(spec.get("title"))
    width = int(spec.get("w") or 420)
    height = int(spec.get("h") or 260)
    chart.width = max(2, width) / 96 * 2.54
    chart.height = max(2, height) / 96 * 2.54
    chart.anchor = _drawing_anchor(sheet, int(spec.get("x") or 0), int(spec.get("y") or 0), width, height)
    ws.add_chart(chart)
    return chart


def _image(sheet: dict, spec: dict, ws) -> None:
    data = str(spec.get("data") or "")
    if "," not in data:
        return
    try:
        raw = base64.b64decode(data.split(",", 1)[1])
        image = XLImage(io.BytesIO(raw))
    except Exception:
        return
    width = int(spec.get("w") or 160)
    height = int(spec.get("h") or 120)
    image.width = width
    image.height = height
    image.anchor = _drawing_anchor(sheet, int(spec.get("x") or 0), int(spec.get("y") or 0), width, height)
    ws.add_image(image)


def _validation(ws, spec: dict) -> None:
    kind = spec.get("type") or "custom"
    if kind not in ("list", "whole", "decimal", "date", "time", "textLength", "custom"):
        kind = "custom"
    operator = spec.get("operator") or "between"
    if operator not in (
        "between",
        "notBetween",
        "equal",
        "notEqual",
        "greaterThan",
        "lessThan",
        "greaterThanOrEqual",
        "lessThanOrEqual",
    ):
        operator = "between"
    formula1 = str(spec.get("formula1") or "")
    formula2 = str(spec.get("formula2") or "")
    if kind == "list" and not formula1.startswith('"') and "," in formula1:
        formula1 = '"%s"' % formula1
    validation = DataValidation(
        type=kind,
        operator=operator,
        formula1=formula1 or None,
        formula2=formula2 or None,
        allow_blank=bool(spec.get("allowBlank", True)),
        showErrorMessage=bool(spec.get("error")),
        showInputMessage=bool(spec.get("prompt")),
        prompt=spec.get("prompt") or None,
        error=spec.get("error") or None,
        errorTitle="mTails",
        promptTitle="mTails",
    )
    ws.add_data_validation(validation)
    for part in str(spec.get("range") or "").split():
        if ":" not in part and not part:
            continue
        try:
            validation.add(part)
        except Exception:
            continue


def _conditional(ws, spec: dict) -> None:
    target = str(spec.get("range") or "")
    if not target:
        return
    kind = spec.get("type")
    if kind == "cellIs":
        fill = None
        color = spec.get("fill")
        if color:
            fill = PatternFill(patternType="solid", fgColor=color.lstrip("#"))
        font = None
        if spec.get("color") or spec.get("bold"):
            font = Font(
                bold=bool(spec.get("bold")),
                color=(spec.get("color") or "000000").lstrip("#"),
            )
        rule = CellIsRule(
            operator=spec.get("operator") or "greaterThan",
            formula=[str(spec.get("formula1") or "0")],
            fill=fill,
            font=font,
        )
        ws.conditional_formatting.add(target, rule)
    elif kind == "colorScale":
        colors = [str(color).lstrip("#") for color in (spec.get("colors") or [])]
        if len(colors) >= 2:
            rule = ColorScaleRule(start_type="min", start_color=colors[0], end_type="max", end_color=colors[-1])
            ws.conditional_formatting.add(target, rule)
    elif kind == "formula":
        ws.conditional_formatting.add(
            target,
            FormulaRule(formula=[str(spec.get("formula1") or "FALSE")], stopIfTrue=False),
        )


def write_xlsx(model: dict) -> bytes:
    """Build an Excel workbook from the model and return it as bytes."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    # Formulas are written without cached results, so ask Excel to recalculate
    # the whole workbook as soon as it is opened.
    workbook.calculation.fullCalcOnLoad = True
    style_cache: Dict[str, Any] = {}

    for index, sheet in enumerate(model.get("sheets") or []):
        name = str(sheet.get("name") or "Sheet%d" % (index + 1))[:31] or "Sheet%d" % (index + 1)
        # Excel forbids these characters and duplicate names.
        for char in "\\/*?:[]":
            name = name.replace(char, " ")
        base = name
        suffix = 1
        while name in workbook.sheetnames:
            suffix += 1
            name = "%s %d" % (base[:28], suffix)
        ws = workbook.create_sheet(name)
        tab_color = sheet.get("tabColor")
        if tab_color:
            try:
                ws.sheet_properties.tabColor = str(tab_color).lstrip("#")
            except Exception:
                pass

        for address, data in (sheet.get("cells") or {}).items():
            if not isinstance(data, dict):
                continue
            formula = data.get("f")
            value = data.get("v")
            if formula:
                text = str(formula)
                cell_value = text if text.startswith("=") else "=" + text
            else:
                cell_value = _value(value)
            try:
                cell = ws[address]
            except Exception:
                continue
            if cell_value is not None:
                cell.value = cell_value
            _apply_style(cell, data.get("s") or {}, style_cache)
            link = data.get("h")
            if link:
                try:
                    cell.hyperlink = str(link)
                    if not data.get("s", {}).get("color"):
                        cell.font = Font(
                            name=(data.get("s") or {}).get("font") or FONT_FALLBACK,
                            size=float((data.get("s") or {}).get("size") or 11),
                            color="0563C1",
                            underline="single",
                        )
                except Exception:
                    pass

        comments = sheet.get("comments") or []
        if isinstance(comments, dict):
            comments = [dict(value, addr=key) for key, value in comments.items() if isinstance(value, dict)]
        for comment in comments:
            if not isinstance(comment, dict):
                continue
            address = comment.get("addr")
            if not address:
                continue
            try:
                ws[address].comment = Comment(str(comment.get("text") or ""), str(comment.get("author") or "mTails"))
            except Exception:
                continue

        for merged in sheet.get("merges") or []:
            try:
                ws.merge_cells(str(merged))
            except Exception:
                continue

        for col, spec in (sheet.get("cols") or {}).items():
            try:
                width = int(spec.get("width") or 96)
                ws.column_dimensions[col].width = max(4.0, (width - 5) / 7.0)
                if spec.get("hidden"):
                    ws.column_dimensions[col].hidden = True
            except Exception:
                continue

        for row, spec in (sheet.get("rows") or {}).items():
            try:
                height = int(spec.get("height") or 22)
                ws.row_dimensions[int(row)].height = max(6.0, height * 72.0 / 96.0)
                if spec.get("hidden"):
                    ws.row_dimensions[int(row)].hidden = True
            except Exception:
                continue

        freeze = sheet.get("freeze") or {}
        rows = int(freeze.get("rows") or 0)
        cols = int(freeze.get("cols") or 0)
        if rows or cols:
            ws.freeze_panes = "%s%d" % (get_column_letter(cols + 1), rows + 1)

        filters = sheet.get("filters")
        if filters and filters.get("range"):
            ws.auto_filter.ref = str(filters["range"])

        for spec in sheet.get("validations") or []:
            _validation(ws, spec)
        for spec in sheet.get("conditional") or []:
            _conditional(ws, spec)
        for spec in sheet.get("charts") or []:
            _chart(sheet, spec, ws)
        for spec in sheet.get("images") or []:
            _image(sheet, spec, ws)

        page = {**default_page(), **(sheet.get("page") or {})}
        ws.page_setup.orientation = "landscape" if page.get("orientation") == "landscape" else "portrait"
        ws.page_setup.paperSize = "1" if page.get("size") == "Letter" else "9"
        try:
            ws.page_setup.scale = max(10, min(400, int(page.get("scale") or 100)))
        except Exception:
            ws.page_setup.scale = 100
        margins = page.get("margins") or {}
        ws.page_margins.left = max(0, float(margins.get("l", 72))) / 96
        ws.page_margins.right = max(0, float(margins.get("r", 72))) / 96
        ws.page_margins.top = max(0, float(margins.get("t", 72))) / 96
        ws.page_margins.bottom = max(0, float(margins.get("b", 72))) / 96
        ws.page_margins.header = max(0, float(margins.get("header", 30))) / 96
        ws.page_margins.footer = max(0, float(margins.get("footer", 30))) / 96
        ws.print_options.gridLines = bool(page.get("gridlines"))
        ws.print_options.headings = bool(page.get("headings"))
        if page.get("header"):
            ws.oddHeader.center.text = str(page["header"])
        if page.get("footer"):
            ws.oddFooter.center.text = str(page["footer"])
        area = page.get("area")
        if area:
            try:
                ws.print_area = str(area)
            except Exception:
                pass

        protection = sheet.get("protection") or {}
        if protection.get("protected"):
            ws.protection = SheetProtection(sheet=True, formatCells=True, formatColumns=True, formatRows=True)

    try:
        active = int(model.get("activeSheet") or 0)
        if 0 <= active < len(workbook.sheetnames):
            workbook.active = active
    except Exception:
        pass

    for name, ref in (model.get("names") or {}).items():
        try:
            workbook.defined_names[str(name)] = DefinedName(str(name), attr_text=str(ref))
        except Exception:
            continue

    try:
        props = model.get("properties") or {}
        workbook.properties.title = str(props.get("title") or model.get("title") or "Workbook1")
        workbook.properties.creator = str(props.get("author") or "mTails")
        workbook.properties.subject = str(props.get("subject") or "")
        workbook.properties.keywords = str(props.get("keywords") or "")
    except Exception:
        pass

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
