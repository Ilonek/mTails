"""Read an Excel workbook into the mTails JSON model (openpyxl)."""

from __future__ import annotations

import base64
import datetime as _dt
import re
from typing import Any, Dict, Optional

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

from .formats import default_page

ERR_VALUES = {
    "#DIV/0!",
    "#VALUE!",
    "#REF!",
    "#NAME?",
    "#N/A",
    "#NUM!",
    "#NULL!",
    "#CIRC!",
    "#GETTING_DATA",
}

BORDER_MAP = {
    "thin": "thin",
    "medium": "medium",
    "thick": "thick",
    "dashed": "dashed",
    "dotted": "dotted",
    "double": "double",
    "hair": "hair",
    "mediumDashed": "mediumDashed",
    "dashDot": "dashDot",
    "mediumDashDot": "mediumDashDot",
    "dashDotDot": "dashDotDot",
}


def _rgb(color) -> Optional[str]:
    if color is None:
        return None
    try:
        value = color.rgb
    except AttributeError:
        return None
    if not isinstance(value, str):
        return None
    value = value.strip()
    if len(value) == 8:
        # openpyxl stores ARGB; the alpha channel carries no meaning for us.
        value = value[2:]
    if len(value) != 6 or not re.match(r"^[0-9A-Fa-f]{6}$", value):
        return None
    return "#" + value.upper()


def _number(value) -> Any:
    if value is None:
        return None
    if isinstance(value, _dt.datetime):
        try:
            return float(to_excel(value))
        except Exception:
            return value.isoformat(sep=" ")
    if isinstance(value, _dt.date):
        try:
            return float(to_excel(_dt.datetime.combine(value, _dt.time())))
        except Exception:
            return value.isoformat()
    if isinstance(value, _dt.time):
        return (value.hour * 3600 + value.minute * 60 + value.second) / 86400.0
    if isinstance(value, _dt.timedelta):
        return value.total_seconds() / 86400.0
    if isinstance(value, float) and value != value:  # NaN
        return None
    try:
        from decimal import Decimal

        if isinstance(value, Decimal):
            return float(value)
    except Exception:
        pass
    if isinstance(value, (int, float, bool, str)):
        return value
    return str(value)


def _style(cell) -> dict:
    style: dict = {}
    font = cell.font
    if font is not None:
        if font.bold:
            style["b"] = True
        if font.italic:
            style["i"] = True
        if font.underline:
            style["u"] = "double" if font.underline == "double" else True
        if font.strike:
            style["strike"] = True
        if font.name and font.name != "Calibri":
            style["font"] = font.name
        if font.sz and float(font.sz) != 11:
            style["size"] = float(font.sz)
        color = _rgb(font.color)
        if color and color != "#000000":
            style["color"] = color
    fill = cell.fill
    if fill is not None and getattr(fill, "patternType", None) == "solid":
        color = _rgb(fill.fgColor)
        if color and color != "#FFFFFF":
            style["fill"] = color
    border = cell.border
    if border is not None:
        sides = {}
        for key, side in (("t", border.top), ("r", border.right), ("b", border.bottom), ("l", border.left)):
            if side is not None and side.style:
                entry = {"style": BORDER_MAP.get(side.style, "thin")}
                color = _rgb(side.color)
                if color:
                    entry["color"] = color
                sides[key] = entry
        if sides:
            style["border"] = sides
    alignment = cell.alignment
    if alignment is not None:
        if alignment.horizontal in ("left", "center", "right", "justify"):
            style["align"] = alignment.horizontal
        if alignment.vertical in ("top", "center", "bottom"):
            style["valign"] = "middle" if alignment.vertical == "center" else alignment.vertical
        if alignment.wrap_text:
            style["wrap"] = True
        if alignment.indent:
            style["indent"] = int(alignment.indent)
    number_format = cell.number_format
    if number_format and number_format != "General":
        style["numFmt"] = number_format
    return style


def _sheet_bounds(sheet: dict) -> tuple[int, int]:
    max_row = 0
    max_col = 0
    for address in sheet.get("cells", {}):
        match = re.match(r"^([A-Z]+)(\d+)$", address)
        if not match:
            continue
        col = 0
        for char in match.group(1):
            col = col * 26 + (ord(char) - 64)
        max_col = max(max_col, col)
        max_row = max(max_row, int(match.group(2)))
    return max_row, max_col


def _col_width(sheet: dict, index: int) -> int:
    from openpyxl.utils import get_column_letter

    entry = sheet.get("cols", {}).get(get_column_letter(index)) or {}
    return int(entry.get("width") or 96)


def _row_height(sheet: dict, index: int) -> int:
    entry = sheet.get("rows", {}).get(str(index)) or {}
    return int(entry.get("height") or 22)


def _pixel_position(sheet: dict, col: int, row: int, x_emu: int, y_emu: int) -> tuple[int, int]:
    """Convert a zero-based openpyxl anchor and EMU offsets into canvas pixels."""
    x = sum(_col_width(sheet, index) for index in range(1, col + 1)) + int(x_emu / 9525)
    y = sum(_row_height(sheet, index) for index in range(1, row + 1)) + int(y_emu / 9525)
    return x, y


def _anchor(chart_or_image) -> tuple[int, int, int, int]:
    anchor = getattr(chart_or_image, "anchor", None)
    start = getattr(anchor, "_from", None)
    col = getattr(start, "col", 0) or 0
    row = getattr(start, "row", 0) or 0
    x_emu = getattr(start, "colOff", 0) or 0
    y_emu = getattr(start, "rowOff", 0) or 0
    return col, row, x_emu, y_emu


def _image_bytes(image) -> Optional[bytes]:
    try:
        data = image._data()
        if isinstance(data, bytes):
            return data
    except Exception:
        pass
    try:
        ref = image.ref
        if hasattr(ref, "seek"):
            ref.seek(0)
        data = ref.read()
        return data if isinstance(data, bytes) else None
    except Exception:
        return None


def _image_mime(data: bytes) -> str:
    if data.startswith(b"\x89PNG"):
        return "image/png"
    if data.startswith(b"\xff\xd8"):
        return "image/jpeg"
    if data.startswith(b"GIF8"):
        return "image/gif"
    if data.startswith(b"BM"):
        return "image/bmp"
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp"
    return "image/png"


def _chart_type(chart) -> str:
    name = chart.__class__.__name__
    if "Bar" in name:
        direction = getattr(chart, "type", "col") or "col"
        return "bar" if direction == "bar" else "column"
    if "Line" in name:
        return "line"
    if "Pie" in name or "Doughnut" in name:
        return "pie"
    if "Area" in name:
        return "area"
    if "Scatter" in name:
        return "scatter"
    return "column"


def _title_text(chart) -> str:
    title = getattr(chart, "title", None)
    if title is None:
        return ""
    try:
        rich = title.tx.rich
        paragraphs = getattr(rich, "p", []) or []
        parts = []
        for paragraph in paragraphs:
            for run in getattr(paragraph, "r", []) or []:
                if getattr(run, "t", None):
                    parts.append(str(run.t))
        return "".join(parts)
    except Exception:
        return ""


def _ref_range(*refs: Optional[str]) -> str:
    """Combine chart references into an A1:B10 style range."""
    cells = []
    for ref in refs:
        if not ref:
            continue
        text = ref.split("!")[-1].replace("$", "")
        if ":" in text:
            cells.extend(text.split(":"))
        else:
            cells.append(text)
    rows = []
    cols = []
    for address in cells:
        match = re.match(r"^([A-Z]+)(\d+)$", address)
        if not match:
            continue
        rows.append(int(match.group(2)))
        col = 0
        for char in match.group(1):
            col = col * 26 + (ord(char) - 64)
        cols.append(col)
    if not rows or not cols:
        return ""
    from openpyxl.utils import get_column_letter

    return "%s%d:%s%d" % (get_column_letter(min(cols)), min(rows), get_column_letter(max(cols)), max(rows))


def _series_refs(series) -> tuple[Optional[str], Optional[str]]:
    val_ref = None
    cat_ref = None
    try:
        if series.val is not None:
            val_ref = (series.val.numRef.f if series.val.numRef else None) or (
                series.val.strRef.f if series.val.strRef else None
            )
    except Exception:
        pass
    for attr in ("cat", "xVal"):
        try:
            cat = getattr(series, attr, None)
            if cat is not None:
                cat_ref = (cat.strRef.f if cat.strRef else None) or (cat.numRef.f if cat.numRef else None)
                if cat_ref:
                    break
        except Exception:
            pass
    return val_ref, cat_ref


def _chart_range(chart) -> str:
    series = getattr(chart, "series", []) or []
    refs = []
    cat_ref = None
    for entry in series:
        val_ref, entry_cat = _series_refs(entry)
        if val_ref:
            refs.append(val_ref)
        if entry_cat and not cat_ref:
            cat_ref = entry_cat
    if not refs:
        return ""
    return _ref_range(*refs, cat_ref)


def _anchor_size(obj) -> tuple[int, int]:
    """Return the pixel size stored in the drawing anchor (or 0, 0)."""
    ext = getattr(getattr(obj, "anchor", None), "ext", None)
    if ext is None:
        return 0, 0
    try:
        return max(0, int(ext.cx / 9525)), max(0, int(ext.cy / 9525))
    except Exception:
        return 0, 0


def _anchor_box(obj, sheet: dict) -> tuple[int, int, int, int]:
    """Return the pixel position and size of a drawing anchor."""
    col, row, x_emu, y_emu = _anchor(obj)
    x, y = _pixel_position(sheet, col, row, x_emu, y_emu)
    anchor = getattr(obj, "anchor", None)
    to = getattr(anchor, "to", None)
    if to is not None:
        x2, y2 = _pixel_position(sheet, to.col or 0, to.row or 0, to.colOff or 0, to.rowOff or 0)
        return x, y, max(24, x2 - x), max(24, y2 - y)
    width, height = _anchor_size(obj)
    return x, y, width, height


def _charts(ws, sheet: dict) -> list:
    items = []
    for chart in getattr(ws, "_charts", []) or []:
        x, y, width, height = _anchor_box(chart, sheet)
        if width < 40 or height < 30:
            # Fall back to the chart's own size, reported in centimetres.
            width = int(float(getattr(chart, "width", 15) or 15) / 2.54 * 96)
            height = int(float(getattr(chart, "height", 7.5) or 7.5) / 2.54 * 96)
        items.append(
            {
                "id": "chart-%d" % (len(items) + 1),
                "type": _chart_type(chart),
                "range": _chart_range(chart),
                "title": _title_text(chart),
                "x": x,
                "y": y,
                "w": width,
                "h": height,
                "headerRow": False,
                "headerCol": False,
            }
        )
    return items


def _images(ws, sheet: dict) -> list:
    items = []
    for image in getattr(ws, "_images", []) or []:
        data = _image_bytes(image)
        if not data:
            continue
        col, row, x_emu, y_emu = _anchor(image)
        x, y, width, height = _anchor_box(image, sheet)
        if width < 24 or height < 24:
            width = int(getattr(image, "width", 160) or 160)
            height = int(getattr(image, "height", 120) or 120)
        mime = _image_mime(data)
        items.append(
            {
                "id": "image-%d" % (len(items) + 1),
                "data": "data:%s;base64,%s" % (mime, base64.b64encode(data).decode("ascii")),
                "x": x,
                "y": y,
                "w": width,
                "h": height,
            }
        )
    return items


def _validations(ws) -> list:
    items = []
    try:
        entries = ws.data_validations.dataValidation
    except AttributeError:
        return items
    for entry in entries or []:
        kind = {
            "whole": "whole",
            "decimal": "decimal",
            "list": "list",
            "date": "date",
            "time": "time",
            "textLength": "textLength",
            "custom": "custom",
        }.get(entry.type, "custom")
        formula1 = entry.formula1 or ""
        formula2 = entry.formula2 or ""
        for sqref in str(entry.sqref).split():
            items.append(
                {
                    "range": sqref,
                    "type": kind,
                    "operator": entry.operator or "between",
                    "formula1": str(formula1).lstrip("="),
                    "formula2": str(formula2).lstrip("="),
                    "allowBlank": bool(entry.allow_blank),
                    "prompt": entry.prompt or "",
                    "error": entry.error or "",
                }
            )
    return items


def _conditional(ws) -> list:
    items = []
    try:
        formatting = list(ws.conditional_formatting)
    except Exception:
        return items
    for cf in formatting:
        for rule in cf.rules:
            entry = {"range": str(cf.sqref), "type": rule.type}
            if rule.type == "cellIs":
                entry.update(
                    {
                        "operator": rule.operator or "greaterThan",
                        "formula1": str(rule.formula[0]).lstrip("=") if rule.formula else "",
                        "formula2": str(rule.formula[1]).lstrip("=") if len(rule.formula or []) > 1 else "",
                    }
                )
                dxf = getattr(rule, "dxf", None)
                if dxf is not None:
                    try:
                        if dxf.fill and dxf.fill.fgColor:
                            entry["fill"] = _rgb(dxf.fill.fgColor)
                    except Exception:
                        pass
                    try:
                        if dxf.font and dxf.font.color:
                            entry["color"] = _rgb(dxf.font.color)
                    except Exception:
                        pass
                    try:
                        entry["bold"] = bool(dxf.font and dxf.font.bold)
                    except Exception:
                        pass
                items.append(entry)
            elif rule.type == "colorScale":
                try:
                    colors = [_rgb(color) for color in rule.colorScale.color]
                except Exception:
                    continue
                entry["colors"] = [color for color in colors if color]
                if len(entry["colors"]) >= 2:
                    items.append(entry)
    return items


def _page_setup(ws) -> dict:
    page = default_page()
    setup = ws.page_setup
    orientation = getattr(setup, "orientation", None) or "portrait"
    page["orientation"] = "landscape" if orientation == "landscape" else "portrait"
    paper = getattr(setup, "paperSize", None)
    page["size"] = "Letter" if paper == "1" else "A4"
    scale = getattr(setup, "scale", None)
    page["scale"] = int(scale) if scale else 100
    margins = ws.page_margins
    page["margins"] = {
        "t": round((margins.top or 0.75) * 96),
        "r": round((margins.right or 0.7) * 96),
        "b": round((margins.bottom or 0.75) * 96),
        "l": round((margins.left or 0.7) * 96),
        "header": round((margins.header or 0.3) * 96),
        "footer": round((margins.footer or 0.3) * 96),
    }
    options = ws.print_options
    page["gridlines"] = bool(getattr(options, "gridLines", False))
    page["headings"] = bool(getattr(options, "headings", False))
    page["header"] = (ws.oddHeader.center.text or "") if ws.oddHeader and ws.oddHeader.center else ""
    page["footer"] = (ws.oddFooter.center.text or "") if ws.oddFooter and ws.oddFooter.center else ""
    area = getattr(ws, "print_area", None)
    if isinstance(area, (list, tuple)):
        area = area[0] if area else ""
    page["area"] = str(area or "").split("!")[-1].replace("$", "")
    return page


def _read_sheet(formula_ws, value_ws, title: str) -> dict:
    sheet: dict = {
        "name": title,
        "cells": {},
        "cols": {},
        "rows": {},
        "merges": [],
        "freeze": {"rows": 0, "cols": 0},
        "charts": [],
        "images": [],
        "comments": [],
        "filters": None,
        "validations": [],
        "conditional": [],
        "tables": [],
        "page": _page_setup(formula_ws),
        "protection": None,
    }

    for entry in formula_ws.column_dimensions.values():
        if entry is None:
            continue
        width = entry.width
        for index in range(entry.min or 1, (entry.max or entry.min or 1) + 1):
            from openpyxl.utils import get_column_letter

            name = get_column_letter(index)
            data = sheet["cols"].setdefault(name, {})
            if width:
                data["width"] = max(12, int(round(width * 7 + 5)))
            if entry.hidden:
                data["hidden"] = True
            if not data:
                sheet["cols"].pop(name, None)

    for index, entry in formula_ws.row_dimensions.items():
        if entry is None:
            continue
        data = sheet["rows"].setdefault(str(index), {})
        if entry.height:
            data["height"] = max(8, int(round(float(entry.height) * 96 / 72)))
        if entry.hidden:
            data["hidden"] = True
        if not data:
            sheet["rows"].pop(str(index), None)

    comment_index = 0
    for cell in list(formula_ws._cells.values()):
        address = cell.coordinate
        if not hasattr(cell, "column") or not hasattr(cell, "row"):
            continue
        raw = cell.value
        formula = None
        value = None
        if isinstance(raw, str) and raw.startswith("="):
            formula = raw
            try:
                value = value_ws[address].value
            except Exception:
                value = None
        else:
            value = raw
        value = _number(value)
        style = _style(cell)
        link = None
        if cell.hyperlink is not None:
            link = cell.hyperlink.target or cell.hyperlink.location
        if cell.comment is not None:
            comment_index += 1
            sheet["comments"].append(
                {
                    "id": comment_index,
                    "addr": address,
                    "text": cell.comment.text or "",
                    "author": cell.comment.author or "mTails",
                    "created": 0,
                }
            )
        if value is None and formula is None and not style and not link:
            continue
        data = {}
        if value is not None:
            data["v"] = value
        if formula:
            data["f"] = formula
        if style:
            data["s"] = style
        if link:
            data["h"] = link
        if data:
            sheet["cells"][address] = data

    for merged in formula_ws.merged_cells.ranges:
        sheet["merges"].append(str(merged))

    freeze = formula_ws.freeze_panes
    if freeze:
        match = re.match(r"^([A-Z]+)(\d+)$", str(freeze).replace("$", ""))
        if match:
            col = 0
            for char in match.group(1):
                col = col * 26 + (ord(char) - 64)
            sheet["freeze"] = {"rows": max(0, int(match.group(2)) - 1), "cols": max(0, col - 1)}

    if formula_ws.auto_filter and formula_ws.auto_filter.ref:
        sheet["filters"] = {"range": str(formula_ws.auto_filter.ref), "criteria": [], "hidden": []}

    sheet["validations"] = _validations(formula_ws)
    sheet["conditional"] = _conditional(formula_ws)
    sheet["charts"] = _charts(formula_ws, sheet)
    sheet["images"] = _images(formula_ws, sheet)

    if formula_ws.protection and formula_ws.protection.sheet:
        sheet["protection"] = {"protected": True, "imported": True}

    try:
        color = formula_ws.sheet_properties.tabColor
        if color is not None and color.rgb:
            rgb = color.rgb[-6:]
            if re.match(r"^[0-9A-Fa-f]{6}$", rgb):
                sheet["tabColor"] = "#" + rgb.upper()
    except Exception:
        pass

    return sheet


def read_xlsx(path: str) -> dict:
    """Load an .xlsx/.xlsm file and return the serialisable workbook model."""
    wb_formulas = load_workbook(path, data_only=False, read_only=False)
    try:
        wb_values = load_workbook(path, data_only=True, read_only=False)
    except Exception:
        wb_values = wb_formulas

    sheets = []
    for ws in wb_formulas.worksheets:
        value_ws = wb_values[ws.title] if ws.title in wb_values.sheetnames else ws
        sheets.append(_read_sheet(ws, value_ws, ws.title))

    names = {}
    try:
        for name, definition in wb_formulas.defined_names.items():
            if str(name).startswith("_xlnm."):
                continue
            if getattr(definition, "type", None) == "RANGE" and definition.attr_text:
                names[str(name)] = str(definition.attr_text).replace("$", "")
    except Exception:
        names = {}

    active = 0
    try:
        active = max(0, wb_formulas.index(wb_formulas.active))
    except Exception:
        pass

    title = "Workbook1"
    properties = {"title": "", "author": "", "subject": "", "keywords": ""}
    try:
        if wb_formulas.properties.title:
            title = str(wb_formulas.properties.title)
        properties = {
            "title": str(wb_formulas.properties.title or title),
            "author": str(wb_formulas.properties.creator or ""),
            "subject": str(wb_formulas.properties.subject or ""),
            "keywords": str(wb_formulas.properties.keywords or ""),
        }
    except Exception:
        pass

    return {
        "version": 1,
        "title": title,
        "properties": properties,
        "activeSheet": active,
        "settings": {"calcMode": "auto", "r1c1": False},
        "names": names,
        "sheets": sheets,
    }
